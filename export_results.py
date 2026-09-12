"""按附件3模板生成四个正式结果工作簿。"""

from __future__ import annotations

from copy import copy
from pathlib import Path

import numpy as np
import openpyxl

from model_core import ROOT

TEMPLATE_DIR = ROOT / "附件/附件3"
OUTPUT_DIR = ROOT / "results"


def interpolate_matrix(
    field: np.ndarray, x_grid: np.ndarray, x_query: np.ndarray
) -> np.ndarray:
    """把全部时刻的场量一次性线性插值到指定坐标。"""

    right = np.searchsorted(x_grid, x_query, side="right")
    right = np.clip(right, 1, len(x_grid) - 1)
    left = right - 1
    width = x_grid[right] - x_grid[left]
    weight = (x_query - x_grid[left]) / width
    return field[:, left] * (1.0 - weight) + field[:, right] * weight


def copy_cell_style(source, target) -> None:
    """扩展模板时沿用原有单元格格式。"""

    target._style = copy(source._style)


def write_sheet(ws, times: np.ndarray, headers: list, matrix: np.ndarray) -> None:
    """覆盖模板示意区并扩展为完整结果表。"""

    header_style = ws.cell(1, 2)
    time_style = ws.cell(2, 1)
    value_style = ws.cell(2, 2)

    for column, header in enumerate(headers, start=2):
        cell = ws.cell(1, column, header)
        copy_cell_style(header_style, cell)
        cell.number_format = "0.0"

    for row_index, time in enumerate(times, start=2):
        time_cell = ws.cell(row_index, 1, float(time))
        copy_cell_style(time_style, time_cell)
        time_cell.number_format = "0"
        for column_index, value in enumerate(matrix[row_index - 2], start=2):
            value_cell = ws.cell(row_index, column_index, float(value))
            copy_cell_style(value_style, value_cell)
            value_cell.number_format = "0.0000"

    ws.freeze_panes = "B2"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = max(ws.column_dimensions["A"].width or 0, 25)
    for column in range(2, len(headers) + 2):
        ws.column_dimensions[openpyxl.utils.get_column_letter(column)].width = 11


def clear_metadata(workbook) -> None:
    """清空输出工作簿中可能暴露身份的属性。"""

    workbook.properties.creator = ""
    workbook.properties.lastModifiedBy = ""
    workbook.properties.title = ""
    workbook.properties.subject = ""
    workbook.properties.description = ""
    workbook.properties.keywords = ""


def export_fixed(problem: int) -> None:
    """导出问题1—3的固定半径结果。"""

    data = np.load(OUTPUT_DIR / f"q{problem}_base.npz")
    time_mask = data["t"] > 0.0
    times = data["t"][time_mask]
    positions_cm = np.round(np.arange(0.0, 2.0 + 0.05, 0.1), 1)
    x_query = positions_cm / 2.0
    moisture = interpolate_matrix(data["moisture"][time_mask], data["x"], x_query)

    workbook = openpyxl.load_workbook(TEMPLATE_DIR / f"result{problem}.xlsx")
    clear_metadata(workbook)
    if problem in (1, 2):
        temperature = interpolate_matrix(
            data["temperature"][time_mask], data["x"], x_query
        )
        write_sheet(workbook["温度"], times, positions_cm.tolist(), temperature)
        write_sheet(workbook["水分浓度"], times, positions_cm.tolist(), moisture)
    else:
        write_sheet(workbook.active, times, positions_cm.tolist(), moisture)
    workbook.save(OUTPUT_DIR / f"result{problem}.xlsx")


def export_moving() -> None:
    """导出问题4的移动边界结果，固定位置始终位于药材内部。"""

    data = np.load(OUTPUT_DIR / "q4_base.npz")
    time_mask = data["t"] > 0.0
    times = data["t"][time_mask]
    radius = data["radius"][time_mask]
    minimum_radius_cm = float(radius.min() * 100.0)
    last_fixed = np.floor((minimum_radius_cm - 1.0e-8) * 10.0) / 10.0
    positions_cm = np.round(np.arange(0.0, last_fixed + 0.05, 0.1), 1)

    rows = []
    field = data["moisture"][time_mask]
    for profile, radius_m in zip(field, radius):
        x_query = positions_cm / (radius_m * 100.0)
        fixed_values = np.interp(x_query, data["x"], profile)
        rows.append(np.append(fixed_values, profile[-1]))
    matrix = np.asarray(rows)

    workbook = openpyxl.load_workbook(TEMPLATE_DIR / "result4.xlsx")
    clear_metadata(workbook)
    headers = positions_cm.tolist() + ["药材表面"]
    write_sheet(workbook.active, times, headers, matrix)
    workbook.save(OUTPUT_DIR / "result4.xlsx")


def main() -> None:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    for problem in (1, 2, 3):
        export_fixed(problem)
        print(f"result{problem}.xlsx 已生成")
    export_moving()
    print("result4.xlsx 已生成")


if __name__ == "__main__":
    main()
