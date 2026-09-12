"""核验四个结果工作簿的结构、数值映射和匿名性。"""

from __future__ import annotations

import json

import numpy as np
import openpyxl

from model_core import ROOT


def expected_value(problem: int, sheet_name: str, time: float, header) -> float:
    """从高精度数组复算指定结果单元格。"""

    data = np.load(ROOT / f"results/q{problem}_base.npz")
    time_index = int(np.where(np.isclose(data["t"], time))[0][0])
    field_name = "temperature" if sheet_name == "温度" else "moisture"
    profile = data[field_name][time_index]
    if header == "药材表面":
        return float(profile[-1])
    x_query = float(header) / (float(data["radius"][time_index]) * 100.0)
    return float(np.interp(x_query, data["x"], profile))


def verify_book(problem: int) -> dict:
    """核验单个结果文件。"""

    path = ROOT / f"results/result{problem}.xlsx"
    workbook = openpyxl.load_workbook(path, data_only=False, read_only=True)
    expected_rows = {1: 1801, 2: 10801, 3: 3450, 4: 3067}[problem]
    expected_columns = {1: 22, 2: 22, 3: 22, 4: 14}[problem]
    step = 1 if problem in (1, 2) else 60
    checks = {
        "file": path.name,
        "sheets": workbook.sheetnames,
        "metadata_anonymous": not any(
            [
                workbook.properties.creator,
                workbook.properties.title,
                workbook.properties.subject,
                workbook.properties.keywords,
            ]
        ),
        "sheet_checks": {},
    }
    for ws in workbook.worksheets:
        row_iterator = ws.iter_rows()
        header_cells = next(row_iterator)
        headers = [cell.value for cell in header_cells[1:]]
        formula_count = 0
        error_count = 0
        times = []
        sample_row_numbers = {2, 2 + (ws.max_row - 1) // 2, ws.max_row}
        sample_column_numbers = {2, 2 + len(headers) // 2, ws.max_column}
        sampled_cells = {}
        values_last = []
        number_format = None
        for row_number, row in enumerate(row_iterator, start=2):
            times.append(float(row[0].value))
            if row_number == 2:
                number_format = row[1].number_format
            if row_number == ws.max_row:
                values_last = [float(cell.value) for cell in row[1:]]
            for column_number, cell in enumerate(row[1:], start=2):
                if cell.data_type == "f":
                    formula_count += 1
                if isinstance(cell.value, str) and cell.value.startswith("#"):
                    error_count += 1
                if (
                    row_number in sample_row_numbers
                    and column_number in sample_column_numbers
                ):
                    sampled_cells[(row_number, column_number)] = float(cell.value)

        times = np.asarray(times, dtype=float)
        maximum_mapping_error = 0.0
        for row in sorted(sample_row_numbers):
            for column in sorted(sample_column_numbers):
                actual = sampled_cells[(row, column)]
                target = expected_value(
                    problem,
                    ws.title,
                    float(times[row - 2]),
                    headers[column - 2],
                )
                maximum_mapping_error = max(maximum_mapping_error, abs(actual - target))

        checks["sheet_checks"][ws.title] = {
            "shape_ok": ws.max_row == expected_rows
            and ws.max_column == expected_columns,
            "time_start": float(times[0]),
            "time_end": float(times[-1]),
            "time_step_ok": bool(np.all(np.diff(times) == step)),
            "headers": headers,
            "formula_count": formula_count,
            "error_count": error_count,
            "maximum_mapping_error": maximum_mapping_error,
            "last_row_min": min(values_last),
            "last_row_max": max(values_last),
            "number_format": number_format,
        }
    return checks


def main() -> None:
    report = {f"result{problem}": verify_book(problem) for problem in (1, 2, 3, 4)}
    output = ROOT / "results/xlsx_validation.json"
    output.write_text(
        json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
